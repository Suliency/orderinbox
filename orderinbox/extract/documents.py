"""Document extraction: turn PDF / XLSX / CSV / EML / plain text into
normalized text plus structured tables.

Deterministic-first: tables come out of the file format itself. The LLM only
sees raw text when the format has no table structure (e.g. a text PDF).
"""
from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass, field
from email import message_from_bytes
from email.header import decode_header
from pathlib import Path
from typing import Any


@dataclass
class DocumentContent:
    path: str = ""
    kind: str = "text"                      # pdf | xlsx | csv | eml | text | html
    text: str = ""                          # full normalized text (for LLM / display)
    tables: list[list[list[str]]] = field(default_factory=list)  # raw 2D strings
    subject: str = ""
    sender: str = ""
    date: str = ""
    attachments: list[str] = field(default_factory=list)  # resolved attachment paths
    ocr_used: bool = False

    @property
    def usable(self) -> bool:
        return bool(self.text.strip() or self.tables)


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------

def _extract_pdf(path: str) -> DocumentContent:
    content = DocumentContent(path=path, kind="pdf")
    text_parts: list[str] = []
    tables: list[list[list[str]]] = []
    try:
        import pdfplumber
    except ImportError:
        return _ocr_fallback(path, content)
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            t = page.extract_text() or ""
            text_parts.append(t)
            try:
                for tbl in page.extract_tables() or []:
                    rows = [[(c or "").strip() for c in row] for row in tbl]
                    if any(any(cell for cell in row) for row in rows):
                        tables.append(rows)
            except Exception:
                pass
    content.text = "\n".join(text_parts).strip()
    content.tables = tables
    if len(content.text) < 40:
        # likely a scanned image-only PDF
        ocr_text = _ocr_text(path)
        if ocr_text:
            content.ocr_used = True
            content.text = (content.text + "\n" + ocr_text).strip()
    return content


def _ocr_text(path: str) -> str:
    """Best-effort OCR via tesseract (pdftoppm -> png -> tesseract) when available."""
    import shutil
    import subprocess
    import tempfile

    if not (shutil.which("tesseract") and shutil.which("pdftoppm")):
        return ""
    with tempfile.TemporaryDirectory() as tmp:
        try:
            subprocess.run(
                ["pdftoppm", "-png", "-r", "200", path, str(Path(tmp) / "page")],
                check=True, capture_output=True, timeout=120,
            )
            parts = []
            for img in sorted(Path(tmp).glob("page-*.png")):
                r = subprocess.run(["tesseract", str(img), "stdout"], capture_output=True, timeout=120)
                if r.returncode == 0:
                    parts.append(r.stdout.decode("utf-8", "ignore"))
            return "\n".join(parts).strip()
        except Exception:
            return ""


def _ocr_fallback(path: str, content: DocumentContent) -> DocumentContent:
    t = _ocr_text(path)
    content.ocr_used = bool(t)
    content.text = t
    return content


# --------------------------------------------------------------------------
# Spreadsheet
# --------------------------------------------------------------------------

def _extract_xlsx(path: str) -> DocumentContent:
    content = DocumentContent(path=path, kind="xlsx")
    try:
        from openpyxl import load_workbook
    except ImportError:
        content.kind = "text"
        content.text = Path(path).read_text(errors="ignore")
        return content
    wb = load_workbook(path, read_only=True, data_only=True)
    text_lines: list[str] = []
    for ws in wb.worksheets:
        rows: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            if row is None:
                continue
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                rows.append(cells)
        if rows:
            content.tables.append(rows)
            header = rows[0]
            for r in rows[1:]:
                text_lines.append(" | ".join(c.strip() for c in r))
    content.text = "\n".join(text_lines).strip()
    return content


def _extract_csv(path: str) -> DocumentContent:
    content = DocumentContent(path=path, kind="csv")
    try:
        raw = Path(path).read_text(errors="ignore")
    except OSError:
        return content
    try:
        rows = [row for row in csv.reader(io.StringIO(raw))]
    except csv.Error:
        content.kind = "text"
        content.text = raw
        return content
    if rows:
        content.tables.append([[c.strip() for c in r] for r in rows])
        content.text = "\n".join("\t".join(r) for r in rows)
    return content


# --------------------------------------------------------------------------
# Email
# --------------------------------------------------------------------------

_TAG_RE = re.compile(r"<[^>]+>")


def _decode_header(value: str | None) -> str:
    if not value:
        return ""
    parts = decode_header(value)
    out = []
    for text, charset in parts:
        if isinstance(text, bytes):
            text = text.decode(charset or "utf-8", "ignore")
        out.append(text)
    return " ".join(out).strip()


def _html_to_text(html: str) -> str:
    html = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", html)
    html = re.sub(r"(?i)<br\s*/?>", "\n", html)
    html = re.sub(r"(?i)</(p|div|tr|li|h[1-6])>", "\n", html)
    html = re.sub(r"(?i)</td>", " | ", html)
    text = _TAG_RE.sub("", html)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n", text)
    return text.strip()


def _extract_eml(path: str, spool_dir: Path) -> DocumentContent:
    content = DocumentContent(path=path, kind="eml")
    try:
        raw = Path(path).read_bytes()
        msg = message_from_bytes(raw)
    except Exception as exc:
        content.text = Path(path).read_text(errors="ignore")
        content.kind = "text"
        return content

    content.subject = _decode_header(msg.get("Subject"))
    content.sender = _decode_header(msg.get("From"))
    content.date = msg.get("Date") or ""

    body_parts: list[str] = []
    att_dir = spool_dir / Path(path).stem
    att_dir.mkdir(parents=True, exist_ok=True)

    if msg.is_multipart():
        for part in msg.walk():
            ctype = part.get_content_type()
            filename = part.get_filename()
            filename = _decode_header(filename) if filename else ""
            if part.get_content_disposition() == "attachment" or (filename and ctype not in ("text/plain", "text/html")):
                data = part.get_payload(decode=True)
                if not data:
                    continue
                safe = re.sub(r"[^A-Za-z0-9._-]", "_", filename) or "attachment.bin"
                dest = att_dir / safe
                dest.write_bytes(data)
                content.attachments.append(str(dest))
                continue
            if ctype == "text/plain":
                payload = part.get_payload(decode=True)
                body_parts.append((payload or b"").decode(part.get_content_charset() or "utf-8", "ignore"))
            elif ctype == "text/html":
                payload = part.get_payload(decode=True)
                html = (payload or b"").decode(part.get_content_charset() or "utf-8", "ignore")
                body_parts.append(_html_to_text(html))
    else:
        payload = msg.get_payload(decode=True)
        if payload is not None:
            body_parts.append(payload.decode(msg.get_content_charset() or "utf-8", "ignore"))
        else:
            body_parts.append(msg.get_payload() or "")

    content.text = "\n".join(p.strip() for p in body_parts if p.strip()).strip()
    return content


# --------------------------------------------------------------------------
# dispatch
# --------------------------------------------------------------------------

def extract_document(path: str | Path, spool_dir: Path | None = None) -> DocumentContent:
    path = Path(path)
    spool_dir = spool_dir or path.parent
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(str(path))
    if suffix in (".xlsx", ".xlsm"):
        return _extract_xlsx(str(path))
    if suffix in (".csv", ".tsv"):
        return _extract_csv(str(path))
    if suffix == ".eml":
        return _extract_eml(str(path), Path(spool_dir))
    text = path.read_text(errors="ignore")
    if suffix == ".html":
        return DocumentContent(path=str(path), kind="html", text=_html_to_text(text))
    return DocumentContent(path=str(path), kind="text", text=text)


def order_like_text(text: str) -> bool:
    """Cheap heuristic: does this text look like it contains an order at all?
    Used to skip newsletters/notifications before spending LLM tokens."""
    if len(text.strip()) < 20:
        return False
    t = text.lower()
    signals = 0
    for kw in ("po number", "po no", "purchase order", "p/o", "order date", "quantity",
               "unit price", "each", "per", "total", "sku", "part no", "part number"):
        if kw in t:
            signals += 1
    # any dollar amounts?
    if re.search(r"\$\s?\d", text):
        signals += 2
    # tabular digits (qty 12 / price 3.45 patterns)?
    if re.search(r"\b\d+\s*(cs|case|ea|each|box|pcs|ct)\b", t):
        signals += 1
    return signals >= 3
